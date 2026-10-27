"""
Build presentation slides for AI Final Exam
Stock Market Trading Agent — Emiliano | SoSe26 | 3sem B
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import os

SLIDES_DIR = "slides_output"
OUT_FILE   = "slides_output/Trading_Agent_v5.pptx"

# ── Color palette ──────────────────────────────────────────────────────────────
DARK_BLUE  = RGBColor(0x1A, 0x2B, 0x4A)
MID_BLUE   = RGBColor(0x2C, 0x5F, 0x9E)
ACCENT     = RGBColor(0x27, 0xAE, 0x60)
LIGHT_GRAY = RGBColor(0xF4, 0xF6, 0xF9)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY  = RGBColor(0x2C, 0x3E, 0x50)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width  = SLIDE_W
prs.slide_height = SLIDE_H

BLANK = prs.slide_layouts[6]  # completely blank layout


def add_bg(slide, color=LIGHT_GRAY):
    from pptx.oxml.ns import qn
    from lxml import etree
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def text_box(slide, text, left, top, width, height,
             font_size=18, bold=False, color=DARK_GRAY,
             align=PP_ALIGN.LEFT, wrap=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txBox


def header_bar(slide, title, subtitle=None):
    """Dark blue top bar with title."""
    bar = slide.shapes.add_shape(
        1, 0, 0, SLIDE_W, Inches(1.15)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK_BLUE
    bar.line.fill.background()

    text_box(slide, title,
             Inches(0.35), Inches(0.12), Inches(12), Inches(0.65),
             font_size=30, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
    if subtitle:
        text_box(slide, subtitle,
                 Inches(0.35), Inches(0.72), Inches(12), Inches(0.38),
                 font_size=14, color=RGBColor(0xAA, 0xCC, 0xFF), align=PP_ALIGN.LEFT)


def accent_line(slide, top=Inches(1.15)):
    line = slide.shapes.add_shape(1, 0, top, SLIDE_W, Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def bullet_box(slide, items, left, top, width, height, font_size=16):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(4)
        run = p.add_run()
        run.text = item
        run.font.size = Pt(font_size)
        run.font.color.rgb = DARK_GRAY


def add_image(slide, path, left, top, width=None, height=None):
    if width and height:
        slide.shapes.add_picture(path, left, top, width, height)
    elif width:
        slide.shapes.add_picture(path, left, top, width=width)
    elif height:
        slide.shapes.add_picture(path, left, top, height=height)
    else:
        slide.shapes.add_picture(path, left, top)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Agent Title
# ══════════════════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(BLANK)
add_bg(s1, DARK_BLUE)

# Big centered title
text_box(s1, "Stock Market Trading Agent",
         Inches(0.5), Inches(1.6), Inches(12.3), Inches(1.4),
         font_size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# Green accent line
accent_line(s1, top=Inches(3.05))

# Agent goal
text_box(s1,
         "An autonomous agent that decides when to buy, sell, or hold Apple (AAPL) stock\n"
         "by analyzing technical price indicators — aiming to maximize portfolio returns.",
         Inches(0.8), Inches(3.2), Inches(11.7), Inches(1.4),
         font_size=20, color=RGBColor(0xDD, 0xEE, 0xFF), align=PP_ALIGN.CENTER)

# Student info
text_box(s1, "Emiliano Velasquez  |  SoSe26  |  3sem B",
         Inches(0.5), Inches(6.5), Inches(12.3), Inches(0.7),
         font_size=15, color=RGBColor(0x88, 0xAA, 0xCC), align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — Environment (Russell & Norvig)
# ══════════════════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(BLANK)
add_bg(s2)
header_bar(s2, "Environment Analysis", "Russell & Norvig Environment Model")
accent_line(s2)

rows = [
    ("Observable",              "Partially Observable",
     "Agent sees prices & volume — but NOT all market forces (news, macro, sentiment)"),
    ("Deterministic vs Stochastic", "Stochastic",
     "Market is driven by countless unpredictable external factors"),
    ("Episodic vs Sequential",  "Sequential",
     "Current trades affect portfolio state and all future decisions"),
    ("Static vs Dynamic",       "Dynamic",
     "Market changes continuously, even while the agent is deliberating"),
    ("Discrete vs Continuous",  "Continuous",
     "Prices and volumes take continuous real values"),
    ("Single vs Multi-agent",   "Multi-agent",
     "Millions of traders acting simultaneously influence prices"),
]

COL_W = [Inches(2.8), Inches(2.5), Inches(7.3)]
COL_X = [Inches(0.15), Inches(2.95), Inches(5.45)]
ROW_H = Inches(0.84)
TABLE_TOP = Inches(1.35)

HEADER_BG  = MID_BLUE
CELL_BG    = WHITE
CHOSEN_BG  = ACCENT
ALT_BG     = RGBColor(0xE8, 0xF4, 0xFD)

header_labels = ["Property", "Classification", "Reasoning"]
for c, (label, x, w) in enumerate(zip(header_labels, COL_X, COL_W)):
    cell = s2.shapes.add_shape(1, x, TABLE_TOP, w, Inches(0.48))
    cell.fill.solid()
    cell.fill.fore_color.rgb = HEADER_BG
    cell.line.color.rgb = WHITE
    text_box(s2, label, x + Inches(0.08), TABLE_TOP + Inches(0.06),
             w - Inches(0.12), Inches(0.38),
             font_size=12, bold=True, color=WHITE)

for r, (prop, chosen, reason) in enumerate(rows):
    y = TABLE_TOP + Inches(0.48) + r * ROW_H
    bg = ALT_BG if r % 2 == 0 else WHITE

    for c, (x, w, content) in enumerate(zip(COL_X, COL_W, [prop, chosen, reason])):
        cell_bg = CHOSEN_BG if c == 1 else bg
        cell = s2.shapes.add_shape(1, x, y, w, ROW_H - Inches(0.04))
        cell.fill.solid()
        cell.fill.fore_color.rgb = cell_bg
        cell.line.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        fcolor = WHITE if c == 1 else DARK_GRAY
        fbold  = True if c == 1 else False
        fsize  = 12 if c == 2 else 13
        text_box(s2, content, x + Inches(0.1), y + Inches(0.1),
                 w - Inches(0.15), ROW_H - Inches(0.15),
                 font_size=fsize, bold=fbold, color=fcolor)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — Data
# ══════════════════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(BLANK)
add_bg(s3)
header_bar(s3, "Dataset", "AAPL — Yahoo Finance  |  2019–2024")
accent_line(s3)

left_items = [
    "Source:  Yahoo Finance (via yfinance library — free, open)",
    "Represents:  Apple Inc. daily stock data, 6 years",
    "Size:  1,460 trading days  |  5 raw features (OHLCV)",
    "Agent sensors:  real-time price feed, volume, computed indicators",
    "",
    "Issues found:",
    "  - Near-balanced classes: 53.8% UP days / 46.2% DOWN days",
    "  - High daily noise — market is inherently stochastic",
    "  - No external data (news, macro) — intentional simplification",
]
bullet_box(s3, left_items, Inches(0.3), Inches(1.3), Inches(5.5), Inches(5.8), font_size=14)

img_path = os.path.join(SLIDES_DIR, "01_price_chart.png")
add_image(s3, img_path, Inches(5.9), Inches(1.25), width=Inches(7.1))

# Features table
feat_data = [
    ("Feature",      "Description",                  "Captures"),
    ("RSI",          "Relative Strength Index (14d)", "Overbought / Oversold"),
    ("SMA_ratio",    "SMA20 / SMA50",                 "Trend direction"),
    ("BB_position",  "Bollinger Band position",       "Proximity to support/resistance"),
    ("MACD",         "12d EMA - 26d EMA",             "Momentum shift"),
    ("Daily_return", "Daily % price change",          "Short-term momentum"),
    ("Volume_change","Daily volume % change",         "Market conviction"),
]

t_top = Inches(5.35)
t_left = Inches(0.3)
t_col_w = [Inches(1.5), Inches(2.5), Inches(1.8)]
t_col_x = [t_left, t_left + Inches(1.5), t_left + Inches(4.0)]
t_row_h = Inches(0.26)

for r, row in enumerate(feat_data):
    y = t_top + r * t_row_h
    is_header = r == 0
    for c, (x, w, val) in enumerate(zip(t_col_x, t_col_w, row)):
        bg = HEADER_BG if is_header else (ALT_BG if r % 2 == 0 else WHITE)
        cell = s3.shapes.add_shape(1, x, y, w, t_row_h)
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
        cell.line.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        fc = WHITE if is_header else DARK_GRAY
        text_box(s3, val, x + Inches(0.06), y + Inches(0.02),
                 w - Inches(0.08), t_row_h - Inches(0.04),
                 font_size=10, bold=is_header, color=fc)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — Data Analysis
# ══════════════════════════════════════════════════════════════════════════════
s4 = prs.slides.add_slide(BLANK)
add_bg(s4)
header_bar(s4, "Data Analysis", "Understanding structure, relationships, and distributions")
accent_line(s4)

add_image(s4, os.path.join(SLIDES_DIR, "03_correlation_heatmap.png"),
          Inches(0.2), Inches(1.3), width=Inches(5.8))
add_image(s4, os.path.join(SLIDES_DIR, "02_class_balance_and_returns.png"),
          Inches(6.2), Inches(1.3), width=Inches(6.9))

notes = [
    "Key findings:",
    "- RSI & BB_position are positively correlated (both track price extremes)",
    "- Daily_return has near-zero correlation with Target — market is hard to predict",
    "- Returns follow a near-normal distribution centered at ~0% (efficient market)",
    "- Class balance is healthy: 53.8% UP vs 46.2% DOWN — no resampling needed",
]
bullet_box(s4, notes, Inches(0.3), Inches(6.1), Inches(12.7), Inches(1.35), font_size=13)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — Model
# ══════════════════════════════════════════════════════════════════════════════
s5 = prs.slides.add_slide(BLANK)
add_bg(s5)
header_bar(s5, "Model & Agent Architecture", "Supervised Learning — Binary Classification")
accent_line(s5)

# Agent architecture diagram (text boxes as diagram)
def diagram_box(slide, text, left, top, width, height, bg=MID_BLUE, fg=WHITE, fsize=13, bold=False):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    shape.line.color.rgb = RGBColor(0x99, 0x99, 0x99)
    text_box(slide, text, left + Inches(0.08), top + Inches(0.06),
             width - Inches(0.14), height - Inches(0.1),
             font_size=fsize, bold=bold, color=fg, align=PP_ALIGN.CENTER)

def arrow(slide, x, y, w=Inches(0.5)):
    txt = slide.shapes.add_textbox(x, y, w, Inches(0.3))
    tf = txt.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = ">>>"
    run.font.size = Pt(14)
    run.font.color.rgb = ACCENT
    run.font.bold = True

BOX_H = Inches(0.55)
ROW_Y = Inches(2.0)
diagram_box(s5, "Percepts\n(Price, Volume)", Inches(0.2),  ROW_Y, Inches(2.0), BOX_H, bg=DARK_GRAY)
arrow(s5, Inches(2.25), ROW_Y + Inches(0.12))
diagram_box(s5, "Feature Engineering\n(RSI, SMA, MACD...)", Inches(2.8), ROW_Y, Inches(2.5), BOX_H, bg=MID_BLUE)
arrow(s5, Inches(5.35), ROW_Y + Inches(0.12))
diagram_box(s5, "Random Forest\nClassifier", Inches(5.9), ROW_Y, Inches(2.3), BOX_H, bg=DARK_BLUE)
arrow(s5, Inches(8.25), ROW_Y + Inches(0.12))
diagram_box(s5, "Decision\nBUY / SELL / HOLD", Inches(8.8), ROW_Y, Inches(2.4), BOX_H, bg=ACCENT)
arrow(s5, Inches(11.25), ROW_Y + Inches(0.12))
diagram_box(s5, "Market\nAction", Inches(11.8), ROW_Y, Inches(1.3), BOX_H, bg=DARK_GRAY)

left_col = [
    "Type:  Supervised Learning",
    "",
    "Primary Model:  Random Forest Classifier",
    "  - n_estimators = 200 trees",
    "  - max_depth = 6",
    "  - Handles non-linear relationships",
    "  - Provides feature importance",
    "",
    "Baseline:  Logistic Regression",
    "  - Used to compare vs. random forest",
    "  - Requires feature scaling (StandardScaler)",
]
bullet_box(s5, left_col, Inches(0.3), Inches(2.85), Inches(5.8), Inches(4.0), font_size=14)

right_col = [
    "Train / Test Split:  Temporal (no shuffle!)",
    "  - Train: Mar 2019 -> Oct 2023  (80%)",
    "  - Test:  Nov 2023 -> Dec 2024  (20%)",
    "  - Shuffling would cause data leakage in time series",
    "",
    "Decision Rule:",
    "  - Predict UP  -> BUY  (if no position held)",
    "  - Predict DOWN -> SELL (if position held)",
    "  - Low confidence -> HOLD",
]
bullet_box(s5, right_col, Inches(6.3), Inches(2.85), Inches(6.7), Inches(4.0), font_size=14)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — Results
# ══════════════════════════════════════════════════════════════════════════════
s6 = prs.slides.add_slide(BLANK)
add_bg(s6)
header_bar(s6, "Model Results", "Random Forest vs. Logistic Regression")
accent_line(s6)

add_image(s6, os.path.join(SLIDES_DIR, "04_confusion_matrices.png"),
          Inches(0.2), Inches(1.3), width=Inches(7.8))
add_image(s6, os.path.join(SLIDES_DIR, "05_feature_importance.png"),
          Inches(8.1), Inches(1.3), width=Inches(5.0))
add_image(s6, os.path.join(SLIDES_DIR, "06_accuracy_comparison.png"),
          Inches(8.1), Inches(4.0), width=Inches(5.0))

notes6 = [
    "Random Forest: 51.4%   |   Logistic Regression: 50.0%   |   Random baseline: 50.0%",
    "Why is accuracy low?  The stock market is stochastic — news, macro events, and sentiment are not in the dataset. This is expected, not a failure.",
    "Improvement ideas: add news sentiment analysis, macro indicators (interest rates, VIX), or switch to Reinforcement Learning.",
]
bullet_box(s6, notes6, Inches(0.3), Inches(6.05), Inches(7.6), Inches(1.4), font_size=13)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — Agent Integration (Final)
# ══════════════════════════════════════════════════════════════════════════════
s7 = prs.slides.add_slide(BLANK)
add_bg(s7)
header_bar(s7, "Model in the Agent System", "From learned model to real-world decisions")
accent_line(s7)

col1 = [
    "How the model fits the agent:",
    "",
    "- Acts as the agent's 'market perception module'",
    "- Each trading day: percepts -> features -> prediction -> action",
    "- Model output feeds directly into the agent's decision logic",
    "- Agent can threshold confidence: only act when P(class) > 0.6",
]
bullet_box(s7, col1, Inches(0.3), Inches(1.4), Inches(6.2), Inches(3.5), font_size=15)

col2 = [
    "What the model enables the agent to understand:",
    "",
    "- Whether current technical indicators signal a buying or selling opportunity",
    "- Which features matter most (Daily_return is the top predictor per feature importance)",
    "- The uncertainty of each prediction (probability scores, not just 0/1)",
]
bullet_box(s7, col2, Inches(6.6), Inches(1.4), Inches(6.5), Inches(3.0), font_size=15)

col3 = [
    "Real-world limitations of this model in the agent:",
    "",
    "- 51% accuracy means high uncertainty -> agent must use risk management (stop-loss, position sizing)",
    "- Model trained on past data -> may fail in new market regimes (e.g. 2020 COVID crash)",
    "- No memory of its own past trades -> a real agent would need state tracking",
    "- Missing percepts: news sentiment, earnings reports, macroeconomic data",
    "- Future direction: combine this supervised model with RL so the agent learns from its own P&L",
]
bullet_box(s7, col3, Inches(0.3), Inches(4.7), Inches(12.7), Inches(2.7), font_size=14)


# ── Save ───────────────────────────────────────────────────────────────────────
prs.save(OUT_FILE)
print(f"Saved: {OUT_FILE}")
